#!/usr/bin/env python3
"""
Проверка результатов лабораторной работы №2
"""

import pandas as pd
import openpyxl
import os

def check_results():
    output_dir = "out/lab2"
    
    print("=== ПРОВЕРКА РЕЗУЛЬТАТОВ ЛАБОРАТОРНОЙ РАБОТЫ №2 ===\n")
    
    # 1. Проверяем основные таблицы задания
    print("1. ОСНОВНЫЕ ТАБЛИЦЫ ЗАДАНИЯ:")
    print("-" * 40)
    
    assignment_file = f"{output_dir}/assignment_tables_final.xlsx"
    if os.path.exists(assignment_file):
        wb = openpyxl.load_workbook(assignment_file)
        print(f"✅ Файл создан: {assignment_file}")
        print(f"   Листы: {wb.sheetnames}")
        
        # Показываем таблицу Prophet
        try:
            prophet_df = pd.read_excel(assignment_file, sheet_name='3_2_Prophet')
            print("\n   Таблица Prophet:")
            print(prophet_df.to_string(index=False))
        except:
            print("   ❌ Ошибка чтения таблицы Prophet")
    else:
        print(f"❌ Файл не найден: {assignment_file}")
    
    # 2. Проверяем график декомпозиции
    print(f"\n2. ГРАФИК ДЕКОМПОЗИЦИИ:")
    print("-" * 40)
    
    decomposition_plot = f"{output_dir}/decomposition_plots_chebyshev.png"
    if os.path.exists(decomposition_plot):
        print(f"✅ График создан: {decomposition_plot}")
        print("   Содержит декомпозицию для 25%, 50%, 75%, 90% данных")
        print("   Показывает: исходные данные, тренд, сезонность (многочлены Чебышева)")
    else:
        print(f"❌ График не найден: {decomposition_plot}")
    
    # 3. Проверяем таблицы декомпозиции Prophet
    print(f"\n3. ТАБЛИЦЫ ДЕКОМПОЗИЦИИ PROPHET:")
    print("-" * 40)
    
    prophet_tables_file = f"{output_dir}/prophet_decomposition_tables.xlsx"
    if os.path.exists(prophet_tables_file):
        print(f"✅ Файл создан: {prophet_tables_file}")
        
        try:
            wb = openpyxl.load_workbook(prophet_tables_file)
            print(f"   Листы: {wb.sheetnames}")
            
            # Показываем сводную информацию
            for sheet in wb.sheetnames:
                df = pd.read_excel(prophet_tables_file, sheet_name=sheet)
                print(f"\n   {sheet}:")
                print(f"   Компоненты: {list(df['Компонента'].values)}")
                
        except Exception as e:
            print(f"   ❌ Ошибка чтения: {e}")
    else:
        print(f"❌ Файл не найден: {prophet_tables_file}")
    
    # 4. Проверяем все графики
    print(f"\n4. ВСЕ ГРАФИКИ:")
    print("-" * 40)
    
    graphics = [
        "original_series_clean_series.png",
        "models_comparison_clean_series.png", 
        "confidence_intervals_clean_series.png",
        "decomposition_plots_chebyshev.png"
    ]
    
    for graphic in graphics:
        path = f"{output_dir}/{graphic}"
        if os.path.exists(path):
            print(f"✅ {graphic}")
        else:
            print(f"❌ {graphic} - не найден")
    
    print(f"\n=== ПРОВЕРКА ЗАВЕРШЕНА ===")

if __name__ == "__main__":
    check_results()
